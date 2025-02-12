import os
import curses
import openpyxl
from datetime import date
import streamlit as st
from typing import List, Optional
import pandas as pd

# Constants
EXCEL_FOLDER = "attendance_sheets"
DEFAULT_ROLL_RANGE = (1, 100)

class ExcelAttendanceManager:
    def __init__(self):
        self.ensure_folders_exist()
    
    @staticmethod
    def ensure_folders_exist():
        """Ensure necessary folders exist"""
        if not os.path.exists(EXCEL_FOLDER):
            os.makedirs(EXCEL_FOLDER)
    
    def get_existing_files(self) -> List[str]:
        """Get list of existing Excel files"""
        return [f for f in os.listdir(EXCEL_FOLDER) if f.endswith(('.xlsx', '.xls'))]
    
    def load_excel(self, file_path: str) -> Optional[openpyxl.Workbook]:
        """Load an Excel workbook"""
        try:
            return openpyxl.load_workbook(file_path)
        except Exception as e:
            st.error(f"Error loading file: {e}")
            return None
    
    def preview_attendance(self, sheet: openpyxl.worksheet.worksheet.Worksheet) -> pd.DataFrame:
        """Show preview of attendance data"""
        data = []
        for row in sheet.iter_rows(min_row=2):
            data.append({
                'Roll Number': row[0].value,
                'Name': row[1].value,
                'Attendance': row[2].value if len(row) > 2 else ''
            })
        return pd.DataFrame(data)
    
    def mark_attendance(self, sheet: openpyxl.worksheet.worksheet.Worksheet, 
                       roll_numbers: List[int], status: str = "A") -> None:
        """Mark attendance for specified roll numbers"""
        today = date.today().strftime("%Y-%m-%d")
        
        # Update header if needed
        if sheet.cell(row=1, column=3).value != 'Attendance':
            sheet.cell(row=1, column=3, value='Attendance')
        
        # Mark attendance
        for row in sheet.iter_rows(min_row=2, min_col=1, max_col=3):
            roll_no = row[0].value
            if roll_no in roll_numbers:
                row[2].value = status
            elif status == "A" and row[2].value is None:  # Mark others as present
                row[2].value = "P"
    
    def get_attendance_stats(self, sheet: openpyxl.worksheet.worksheet.Worksheet) -> dict:
        """Get attendance statistics"""
        total = present = absent = 0
        for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
            if row[0].value:
                total += 1
                if row[0].value == 'P':
                    present += 1
                elif row[0].value == 'A':
                    absent += 1
        
        return {
            'total': total,
            'present': present,
            'absent': absent,
            'percentage': round((present/total * 100), 2) if total > 0 else 0
        }

def file_selector_curses(stdscr, files: List[str]) -> Optional[str]:
    """Use curses to select a file"""
    curses.init_pair(1, curses.COLOR_GREEN, curses.COLOR_BLACK)
    curses.init_pair(2, curses.COLOR_BLACK, curses.COLOR_WHITE)
    
    current_row = 0
    
    while True:
        stdscr.clear()
        height, width = stdscr.getmaxyx()
        
        # Show header
        stdscr.addstr(0, 0, "Select Attendance File (Use arrow keys, Enter to select)", curses.A_BOLD)
        
        # Show files
        for idx, file in enumerate(files):
            if idx == current_row:
                stdscr.attron(curses.color_pair(2))
                stdscr.addstr(idx + 2, 0, f"> {file}")
                stdscr.attroff(curses.color_pair(2))
            else:
                stdscr.addstr(idx + 2, 0, f"  {file}")
        
        key = stdscr.getch()
        
        if key == curses.KEY_UP and current_row > 0:
            current_row -= 1
        elif key == curses.KEY_DOWN and current_row < len(files) - 1:
            current_row += 1
        elif key == ord('\n'):  # Enter key
            return files[current_row]
        elif key == ord('q'):  # Quit
            return None
        
        stdscr.refresh()

def main():
    st.title("📚 Attendance Manager")
    
    manager = ExcelAttendanceManager()
    
    # File upload
    uploaded_files = st.file_uploader("Upload Excel files", type=['xlsx', 'xls'], accept_multiple_files=True)
    
    if uploaded_files:
        for uploaded_file in uploaded_files:
            file_path = os.path.join(EXCEL_FOLDER, uploaded_file.name)
            with open(file_path, "wb") as f:
                f.write(uploaded_file.getbuffer())
    
    # Get existing files
    excel_files = manager.get_existing_files()
    
    if not excel_files:
        st.warning("No attendance files found. Please upload an Excel file.")
        return
    
    # File selection
    st.subheader("Select Attendance File")
    selected_file = st.selectbox("Choose File", excel_files)
    
    if selected_file:
        file_path = os.path.join(EXCEL_FOLDER, selected_file)
        workbook = manager.load_excel(file_path)
        
        if workbook:
            # Sheet selection
            sheet_names = workbook.sheetnames
            selected_sheet = st.selectbox("Select Sheet", sheet_names)
            sheet = workbook[selected_sheet]
            
            # Show current attendance
            st.subheader("Current Attendance")
            df = manager.preview_attendance(sheet)
            st.dataframe(df)
            
            # Mark attendance
            st.subheader("Mark Attendance")
            absent_input = st.text_input("Enter absent roll numbers (comma-separated):")
            
            if st.button("Mark Attendance"):
                try:
                    # Process absent numbers
                    absent_rolls = [int(roll.strip()) for roll in absent_input.split(',') if roll.strip()]
                    
                    # Validate roll numbers
                    if not all(DEFAULT_ROLL_RANGE[0] <= roll <= DEFAULT_ROLL_RANGE[1] for roll in absent_rolls):
                        st.error("❌ Invalid roll numbers! Please enter numbers between 1-100.")
                        return
                    
                    # Mark attendance
                    manager.mark_attendance(sheet, absent_rolls, "A")
                    workbook.save(file_path)
                    
                    # Show success and stats
                    st.success("✅ Attendance marked successfully!")
                    stats = manager.get_attendance_stats(sheet)
                    
                    # Show statistics
                    col1, col2, col3 = st.columns(3)
                    with col1:
                        st.metric("Present", stats['present'])
                    with col2:
                        st.metric("Absent", stats['absent'])
                    with col3:
                        st.metric("Attendance %", f"{stats['percentage']}%")
                    
                    # Show the updated attendance sheet directly
                    st.subheader("Updated Attendance Sheet")
                    updated_df = manager.preview_attendance(sheet)
                    st.dataframe(updated_df)
                
                except ValueError as e:
                    st.error(f"❌ Error: {str(e)}")
                except Exception as e:
                    st.error(f"❌ An unexpected error occurred: {str(e)}")

if __name__ == "__main__":
    main()